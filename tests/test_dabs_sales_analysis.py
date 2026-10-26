"""DABS monthly Sales Analysis Report: grouping by beverage, the printed totals, and the report's own period."""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

from alcohol_data import paths
from alcohol_data.adapters.base import Rejected, Staged
from alcohol_data.adapters.dabs_sales_analysis import beverage_map, parse, shelf_price_profile
from alcohol_data.config import AssetConfig
from alcohol_data.transforms.derive import dabs_fiscal_years

# (class code, class name, gross sales, 9-liter cases); ADF spirits, EDD wine, LAB beer, ZZZ supplies.
PRODUCTS = [
    ("ADF", "VODKA - BASIC", 1000.25, 10.5),
    ("ADF", "VODKA - BASIC", 500.00, 4.25),
    ("EDD", "VERMOUTH - DRY", 300.10, 2.0),
    ("LAB", "ALLOCATED BEER", 99.65, 1.125),
    ("ZZZ", "SUPPLIES", 0.0, 0.0),
]


def _asset(fy=2026, period=1, month="2025-07") -> AssetConfig:
    return AssetConfig(
        "dabs_test",
        "utah_dabs_sales_analysis",
        None,
        "manual_import",
        "e",
        "utah_dabs_sales_analysis",
        "v",
        "xlsx",
        "dabs_sales_analysis",
        "t.xlsx",
        options={"fiscal_year": fy, "period": period, "month": month},
    )


def _xlsx(
    tmp_path: Path, products=PRODUCTS, period="FY26 P1", month=datetime(2025, 7, 1), totals=None, col=2
) -> Path:
    """A report shaped like DABS's: title block with the period, the month and two printed totals, then a
    heading row whose 'Class Code' is not in column A."""
    if totals is None:
        totals = (sum(p[2] for p in products), sum(p[3] for p in products))
    wb = openpyxl.Workbook()
    ws = wb.active
    pad = [None] * col
    ws.append([*pad, "Sales Analysis Report", period])
    ws.append([*pad, month, None, "Gross Sales", totals[0], "Case Count (9L) Sales", totals[1]])
    ws.append([])
    head = [
        "Class Code",
        "Class Name",
        "Item",
        "Description",
        "Size",
        "Store",
        "Agency",
        "Total",
        "a",
        "b",
        "c",
        "d",
        "Total",
    ]
    ws.append([*pad, *head])
    for code, name, sales, cases in products:
        ws.append([*pad, code, name, "1", "x", "750ml", None, None, sales, None, None, None, None, cases])
    p = tmp_path / "Sales-Analysis.xlsx"
    wb.save(p)
    return p


def _by_measure(items):
    assert not [i for i in items if isinstance(i, Rejected)], [i for i in items if isinstance(i, Rejected)]
    return {s.measure_id: s for s in items if isinstance(s, Staged)}


def test_groups_products_by_beverage_and_keeps_the_printed_totals(tmp_path):
    got = _by_measure(list(parse(_asset(), _xlsx(tmp_path))))
    assert got["ut_dabs_monthly_gross_sales_spirits"].value == 1500.25
    assert got["ut_dabs_monthly_cases_9l_spirits"].value == 14.75
    assert got["ut_dabs_monthly_gross_sales_wine"].value == 300.10
    assert got["ut_dabs_monthly_gross_sales_beer"].value == 99.65
    assert (
        got["ut_dabs_monthly_gross_sales_cider"].value == 0
    )  # no cider products: a real zero, the report lists every product
    assert got["ut_dabs_monthly_gross_sales_other"].value == 0
    assert got["ut_dabs_monthly_gross_sales_total"].value == pytest.approx(1900.0)
    s = got["ut_dabs_monthly_gross_sales_wine"]
    assert (s.geo_id, s.period_kind, s.period_label, s.period_start, s.period_end) == (
        "state:49",
        "month",
        "2025-07",
        date(2025, 7, 1),
        date(2025, 7, 31),
    )
    assert s.extra["dabs_fiscal_period"] == "FY26 P1"


def test_rows_that_do_not_add_up_to_the_printed_total_reject_the_whole_report(tmp_path):
    items = list(parse(_asset(), _xlsx(tmp_path, totals=(1900.01, 17.875))))
    assert len(items) == 1 and isinstance(items[0], Rejected) and "add up" in items[0].reason


def test_an_unmapped_class_code_rejects_the_report_rather_than_dropping_the_row(tmp_path):
    items = list(parse(_asset(), _xlsx(tmp_path, products=[*PRODUCTS, ("QQQ", "NEW CLASS", 5.0, 1.0)])))
    assert len(items) == 1 and isinstance(items[0], Rejected) and "QQQ" in items[0].reason


@pytest.mark.parametrize("kw", [{"period": "FY26 P2"}, {"month": datetime(2025, 8, 1)}])
def test_a_report_for_another_period_is_rejected(tmp_path, kw):
    items = list(parse(_asset(), _xlsx(tmp_path, **kw)))
    assert len(items) == 1 and isinstance(items[0], Rejected)


def test_every_class_code_maps_to_exactly_one_known_group():
    groups = beverage_map()
    assert set(groups.values()) <= {
        "spirits",
        "wine",
        "beer",
        "cider",
        "premixed",
        "unclassified",
        "non_beverage",
    }
    assert groups["ADF"] == "spirits" and groups["EDD"] == "wine" and groups["ZZZ"] == "non_beverage"


def _month_obs(measure, y, m, value, i):
    return {
        "measure_id": measure,
        "period_kind": "month",
        "geo_id": "state:49",
        "period_start": date(y, m, 1),
        "value": value,
        "observation_id": f"o{i}",
    }


def test_a_fiscal_year_is_summed_only_from_all_twelve_months():
    months = [(2025, m) for m in range(7, 13)] + [(2026, m) for m in range(1, 7)]
    full = [_month_obs("ut_dabs_monthly_gross_sales_wine", y, m, 10.0, i) for i, (y, m) in enumerate(months)]
    ((row, inputs, formula),) = dabs_fiscal_years(full)
    assert row["measure_id"] == "ut_dabs_fy_gross_sales_wine" and row["value"] == 120.0
    assert (row["period_start"], row["period_end"], row["period_label"]) == (
        date(2025, 7, 1),
        date(2026, 6, 30),
        "FY2026",
    )
    assert inputs[0] == "o0" and len(inputs) == 12 and formula.endswith("dabs_fy_sum")
    assert (
        dabs_fiscal_years(full[:-1]) == []
    )  # eleven months: no fiscal-year value at all, never a partial sum
    gap = [dict(o) for o in full]
    gap[3]["value"] = None
    assert dabs_fiscal_years(gap) == []


# ------------------------------------------------------------------ shelf-price basis of gross sales

REAL_HEAD = [
    "Class Code",
    "Class Name",
    "Bottle Size",
    "AG Name",
    "DG Name",
    "Item Code",
    "Item Name",
    "Total",
    "Percent",
    "YOY",
    "Rank",
    "Bottle Count Sales",
    "Total",
    "Percent",
]


def _row(sales, bottles, ml=750, cases=None):
    """A product row in DABS's own column order: gross sales at +7, bottles at +11, 9-liter cases at +12."""
    if cases is None:
        cases = bottles * ml / 9000
    return ["ADF", "VODKA", str(ml), "a", "d", "000001", "X", sales, None, None, None, bottles, cases, None]


def test_shelf_price_profile_tells_shelf_prices_from_tax_folded_totals():
    rows = [
        REAL_HEAD,
        _row(20 * 29.99, 20),  # 20 bottles at a $29.99 shelf price: whole cents
        _row(
            4453 * 24.99 + 12 * 29.99, 4465
        ),  # Aperol, Aug 2026: two shelf prices in the month, not whole cents
        _row(643.29, 20),  # 20 x $29.99 with 7.25% tax folded in (643.2855, to the cent): not whole cents
        _row(5 * 29.99, 5),  # under 10 bottles: not tested
        _row(10 * 9.99, 10, ml=1000, cases=1.5),  # cases disagree with 10 x 1,000 ml (1.111): counted against
    ]
    got = shelf_price_profile(rows, 0, 0)
    assert got == {
        "min_bottles": 10,
        "products_tested": 4,
        "whole_cent_unit_price": 2,
        "gross_sales_tested": round(599.80 + 111640.35 + 643.29 + 99.90, 2),
        "gross_sales_whole_cent": round(599.80 + 99.90, 2),
        "products_with_size": 5,
        "bottles_match_cases": 4,
    }


def test_no_profile_without_the_bottle_column_where_dabs_prints_it():
    head = [*REAL_HEAD]
    head[11] = "Something Else"
    assert shelf_price_profile([head, _row(599.80, 20)], 0, 0) is None
    assert shelf_price_profile([REAL_HEAD[:12], _row(599.80, 20)[:12]], 0, 0) is None


def _real_report(asset_id: str) -> Path | None:
    """The checkout's own raw cache (the tests' data home is a temporary folder)."""
    found = sorted((paths.ROOT / "data" / "raw" / asset_id).glob("*.xlsx"))
    return found[0] if found else None


AUG_2026 = _real_report("utah_dabs_sales_analysis_fy2027_p02")


@pytest.mark.skipif(AUG_2026 is None, reason="DABS August 2026 report not in the raw cache")
def test_august_2026_report_reads_as_shelf_prices_times_bottles():
    """Hand-checkable against DABS's August 2026 price lists (new prices effective 2026-08-01): Crown Royal
    Regal Apple 750ml (010807) $29.99; 1800 Silver Tequila 750ml (087510) new retail $26.99; Aperol 750ml
    (064147) $29.99 cut to $24.99. Checked on 2026-10-11 from abs.utah.gov/wp-content/uploads/
    Aug2026NumericPriceList.pdf and Aug2026AlphaPriceList.pdf (read online, not pinned)."""
    ws = openpyxl.load_workbook(AUG_2026, read_only=True, data_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(rows) if "Class Code" in r)
    c = rows[hi].index("Class Code")
    items = {r[c + 5]: (r[c + 7], r[c + 11]) for r in rows[hi + 1 :] if r and r[c] is not None}
    assert items["010807"] == (pytest.approx(136034.64), 4536) and round(136034.64 / 4536, 2) == 29.99
    assert items["087510"] == (pytest.approx(113223.05), 4195) and round(113223.05 / 4195, 4) == 26.99
    assert items["064147"] == (pytest.approx(111640.35), 4465)
    assert round(4453 * 24.99 + 12 * 29.99, 2) == 111640.35  # old and new shelf price, nothing else
    prof = shelf_price_profile(rows, hi, c)
    assert (prof["products_tested"], prof["whole_cent_unit_price"]) == (5013, 4581)
    assert prof["bottles_match_cases"] == prof["products_with_size"]


def test_shelf_price_check_flags_months_without_a_profile_or_with_few_whole_cent_prices():
    duckdb = pytest.importorskip("duckdb")
    from alcohol_data.validation.checks import r_dabs_shelf

    con = duckdb.connect()
    con.execute(
        "CREATE TABLE observations (observation_id VARCHAR, measure_id VARCHAR, period_label VARCHAR, period_start DATE)"
    )
    con.execute("CREATE TABLE observation_details (observation_id VARCHAR, extra VARCHAR)")

    def prof(tested, whole, sized=100, match=100):
        return json.dumps(
            {
                "dabs_fiscal_period": "FY26 P1",
                "shelf_price_profile": {
                    "products_tested": tested,
                    "whole_cent_unit_price": whole,
                    "products_with_size": sized,
                    "bottles_match_cases": match,
                },
            }
        )

    months = [
        ("a", "2025-07", prof(100, 90)),
        ("b", "2025-08", prof(100, 70)),  # below 75%
        ("c", "2025-09", json.dumps({"dabs_fiscal_period": "FY26 P3"})),  # no profile
        ("d", "2025-10", prof(100, 95, sized=100, match=98)),  # bottle column disagrees with cases
    ]
    for i, (oid, label, extra) in enumerate(months):
        con.execute(
            "INSERT INTO observations VALUES (?, 'ut_dabs_monthly_gross_sales_total', ?, ?)",
            [oid, label, date(2025, 7 + i, 1)],
        )
        con.execute("INSERT INTO observation_details VALUES (?, ?)", [oid, extra])
    details, summary = r_dabs_shelf(con)
    assert [d.split(":")[0] for d in details] == ["2025-08", "2025-09", "2025-10"]
    assert "255 of 300" in summary
