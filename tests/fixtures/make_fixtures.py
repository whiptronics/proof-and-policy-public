"""Regenerate the test fixtures from the pinned raw files (maintainer tool, needs 'alcohol-data fetch' first).

Every fixture is a subset of a real U.S. federal source file (public domain, except the
BRFSS 1995-2010 slice, ODC-BY; see THIRD_PARTY_NOTICES.md). Text fixtures keep the
original lines byte-for-byte; the STC workbook fixture is rebuilt with openpyxl from
real rows of the official workbook (same sheet name, header rows and cell values).

    python tests/fixtures/make_fixtures.py
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw"
OUT = Path(__file__).resolve().parent


def raw(asset_id: str) -> Path:
    files = [p for p in (RAW / asset_id).iterdir() if not p.name.endswith(".meta.json")]
    if len(files) != 1:
        raise SystemExit(f"expected one cached file for {asset_id}, found {files}")
    return files[0]


def lines_matching(path: Path, keep, encoding="latin-1") -> str:
    out = []
    with open(path, encoding=encoding, newline="") as f:
        for line in f:
            if keep(line):
                out.append(line)
    return "".join(out)


def main() -> None:
    # NIAAA: full header (layout block) + UT, TN, West region, U.S. for 2014-2017 and 2024.
    data_line = re.compile(r"^(\d{4}) +(\d+) ")

    def keep_niaaa(line: str) -> bool:
        m = data_line.match(line)
        if not m:
            return True
        return m.group(2) in ("47", "49", "94", "99") and m.group(1) in (
            "2014",
            "2015",
            "2016",
            "2017",
            "2024",
        )

    (OUT / "niaaa_pcyr_subset.txt").write_text(
        lines_matching(raw("niaaa_pcyr_sr124"), keep_niaaa), encoding="latin-1", newline=""
    )

    # Census 2000s single-year-of-age: header + Alabama (published as STATE "1") and Utah, both sexes total.
    p = raw("census_pep_icen_2000s_agesex")
    text = p.read_text(encoding="latin-1")
    rows = text.splitlines(keepends=True)
    keep = [rows[0]] + [r for r in rows[1:] if re.match(r"^\d+,\d+,(1|49),[^,]+,0,", r)]
    (OUT / "census_st_agesex_2000s_subset.csv").write_text("".join(keep), encoding="latin-1", newline="")

    # Census Vintage 2025 alldata6: header + Utah rows for SEX 0 and SEX 1 (SEX 1 must be ignored).
    p = raw("census_pep_v2025_alldata6")
    rows = p.read_text(encoding="latin-1").splitlines(keepends=True)
    keep = [rows[0]] + [r for r in rows[1:] if re.match(r"^040,\d,\d,49,Utah,(0|1),", r)]
    (OUT / "census_alldata6_v2025_utah.csv").write_text("".join(keep), encoding="latin-1", newline="")

    # Census 1990s intercensal: two small Utah counties (Daggett 49009, Piute 49031), 1995.
    p = raw("census_pep_icen_1995")
    sub = lines_matching(p, lambda line: line[4:9] in ("49009", "49031"))
    (OUT / "census_icen1995_two_utah_counties.txt").write_text(sub, encoding="latin-1", newline="")

    # STC FY2025 flat file, complete (8 KB).
    (OUT / "stc_fy2025_flat.txt").write_bytes(raw("census_stc_fy2025_flat").read_bytes())

    # STC historical workbook subset: header + description rows, then real rows for selected years/states.
    wb = openpyxl.load_workbook(raw("census_stc_historical_db_2025"), read_only=True)
    ws = wb.worksheets[0]
    src = list(ws.iter_rows(values_only=True))
    keep_rows = [
        r
        for r in src[2:]
        if r[0] is not None
        and (
            (r[0] in (2014, 2015, 2016, 2017, 2025) and r[1] in (0, 1, 43, 45))
            or (r[0] == 1955 and r[1] == 1)  # fiscal year ends in 1954 (before the STC year)
            or (r[0] == 1922 and r[1] == 0)
        )
    ]  # national row with the -11111 sentinel
    nb = openpyxl.Workbook()
    ns = nb.active
    ns.title = ws.title
    for r in src[:2] + keep_rows:
        ns.append(list(r))
    nb.save(OUT / "stc_historical_subset.xlsx")

    # BRFSS: Utah, Tennessee, plus median and territory rows (which must be skipped), selected years.
    for asset, fname, years in (
        (
            "cdc_brfss_2011_present_alcohol_overall",
            "brfss_2011_present_subset.csv",
            ("2011", "2015", "2022", "2024"),
        ),
        (
            "cdc_brfss_1995_2010_alcohol_overall",
            "brfss_1995_2010_subset.csv",
            ("1995", "2005", "2006", "2010"),
        ),
    ):
        text = raw(asset).read_text(encoding="utf-8")
        reader = csv.reader(io.StringIO(text))
        header = next(reader)
        yi, li = header.index("year"), header.index("locationabbr")
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\n")
        w.writerow(header)
        for r in reader:
            if r[yi] in years and r[li] in ("UT", "TN", "UW", "PR"):
                w.writerow(r)
        (OUT / fname).write_text(buf.getvalue(), encoding="utf-8", newline="")

    # BLS: the 2017-2026 response, reduced to 2023-2025 (includes the unavailable October 2025 value).
    payload = json.loads(raw("bls_cpi_u_2017_2026").read_text(encoding="utf-8"))
    s = payload["Results"]["series"][0]
    s["data"] = [d for d in s["data"] if d["year"] in ("2023", "2024", "2025")]
    (OUT / "bls_cpi_2023_2025.json").write_text(
        json.dumps(payload, indent=1) + "\n", encoding="utf-8", newline="\n"
    )
    print("fixtures written to", OUT)


if __name__ == "__main__":
    main()
