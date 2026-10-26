"""Census national gross margins as a percent of sales: Annual Retail Trade Survey (gmper.xlsx) and Annual Wholesale
Trade Survey table 4 (merchant wholesalers except manufacturers' sales branches and offices).

Only the industries the grocery-wine comparison refers to are read: grocery stores, food and beverage stores, beer,
wine and liquor stores, warehouse clubs and supercenters and the retail total; beer, wine and distilled alcoholic
beverage wholesalers, grocery wholesalers and the merchant-wholesale total. Every row is matched by its NAICS code
and checked against its printed label, so a re-laid-out sheet is rejected rather than misread. Years come from the
header row ("2022r" is a revised year, flagged "revised"). "S" (does not meet publication standards) is stored as
suppressed, the wholesale table's "x" (no prior-year inventories, so no 1992 margin) and "NA" as unavailable, and a
blank cell as not published; none is ever zero. National, all products, calendar data years.
A gross margin is sales less cost of goods sold, as a percent of sales: not a markup, not profit.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from alcohol_data.adapters.base import NATION, Rejected, Staged, register
from alcohol_data.config import AssetConfig

# NAICS code (or TOTAL) -> (measure, the label the row must start with)
ARTS = {
    "TOTAL": ("us_retail_gross_margin_pct_total", "Retail gross margin as a percentage of sales, total"),
    "445": ("us_retail_gross_margin_pct_food_beverage_stores", "Food and beverage stores"),
    "4451": ("us_retail_gross_margin_pct_grocery_stores", "Grocery stores"),
    "4453": ("us_retail_gross_margin_pct_alcohol_stores", "Beer, wine, and liquor stores"),
    "452311": ("us_retail_gross_margin_pct_warehouse_clubs_supercenters", "Warehouse clubs and supercenters"),
}
AWTS = {
    "42": ("us_wholesale_gross_margin_pct_total", "Wholesale Trade"),
    "4244": ("us_wholesale_gross_margin_pct_grocery", "Grocery and Related Products"),
    "4248": (
        "us_wholesale_gross_margin_pct_alcohol_beverages",
        "Beer, Wine, and Distilled Alcoholic Beverages",
    ),
}
AWTS_ITEM = "Gross margins as a percent of sales"
AWTS_OPERATION = "Merchant Wholesalers, except manufacturers' sales branches and offices"
YEAR = re.compile(r"^(\d{4})(r?)$")
# The wholesale table's legend; the retail table uses only "S".
UNAVAILABLE = {
    "x": "Estimate is not available due to the lack of prior year inventories.",
    "NA": "Estimate is unavailable.",
}
# Leader dots and the replacement characters the sheet's label column carries after each label.
LEADER = re.compile(r"[\s.…�]+$")


def _label(v) -> str:
    return LEADER.sub("", str(v or "")).strip()


def _years(row: tuple) -> list[tuple[int, int, bool, str]]:
    """(column index, year, revised, header text) for every year-like header cell."""
    out = []
    for i, v in enumerate(row):
        m = YEAR.match(str(v).strip()) if v is not None else None
        if m:
            out.append((i, int(m.group(1)), m.group(2) == "r", str(v).strip()))
    return out


def _cells(
    asset: AssetConfig,
    path: Path,
    sheet: str,
    row_no: int,
    row: tuple,
    years,
    measure: str,
    code: str,
    label: str,
    series: str,
):
    for col, y, revised, head in years:
        raw = row[col] if col < len(row) else None
        coord = f"{get_column_letter(col + 1)}{row_no}"
        base = dict(
            geo_id=NATION,
            measure_id=measure,
            period_start=date(y, 1, 1),
            period_end=date(y, 12, 31),
            period_kind="calendar_year",
            period_label=str(y),
            year=y,
            estimate_type="survey_estimate",
            extraction_method="spreadsheet_parser",
            comparability_group=f"{series}|{measure}",
            source_code=f"NAICS {code}, {label}, {head}",
            raw_unit="percent of sales",
            multiplier=1.0,
            source_locator=f"{path.name} sheet '{sheet}' cell {coord} (NAICS {code}, {head})",
            quality_flags=["revised"] if revised else [],
        )
        text = "" if raw is None else str(raw).strip()
        if text == "":
            yield Staged(value=None, missing_reason="not_published", raw_value=text, **base)
        elif text == "S":
            yield Staged(
                value=None, missing_reason="suppressed", suppression_flag="S", raw_value=text, **base
            )
        elif text in UNAVAILABLE:
            yield Staged(
                value=None, missing_reason="unavailable", raw_value=text, footnote=UNAVAILABLE[text], **base
            )
        elif isinstance(raw, (int, float)):
            yield Staged(value=float(raw), raw_value=text, **base)
        else:
            yield Rejected(
                base["source_locator"],
                f"unreadable value {text!r}",
                measure_id=measure,
                geo_id=NATION,
                period_label=str(y),
                raw_value=text,
            )


@register("census_arts_gmper", "1.0")
def parse_arts(asset: AssetConfig, path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head_at = next((i for i, r in enumerate(rows) if r and str(r[0] or "").strip() == "NAICS Code"), None)
    if head_at is None:
        yield Rejected(path.name, "no 'NAICS Code' header row")
        return
    years = _years(rows[head_at])
    if not years:
        yield Rejected(path.name, "no year columns in the header row")
        return
    found: set[str] = set()
    for i, r in enumerate(rows[head_at + 1 :], start=head_at + 2):
        if not r or len(r) < 2:
            continue
        label = _label(r[1])
        code = str(r[0]).strip() if r[0] is not None else ""
        if not code and label == ARTS["TOTAL"][1]:
            code = "TOTAL"
        if code not in ARTS:
            continue
        measure, want = ARTS[code]
        if not label.startswith(want) or code in found:
            yield Rejected(
                f"{path.name} row {i}",
                f"NAICS {code}: label {label!r}, expected {want!r}, or a second row",
                measure_id=measure,
            )
            continue
        found.add(code)
        yield from _cells(asset, path, ws.title, i, r, years, measure, code, label, asset.source_series)
    for code in ARTS.keys() - found:
        yield Rejected(path.name, f"NAICS {code} ({ARTS[code][1]}) not found", measure_id=ARTS[code][0])


@register("census_awts_table4", "1.0")
def parse_awts(asset: AssetConfig, path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head_at = next(
        (
            i
            for i, r in enumerate(rows)
            if r and "NAICS" in str(r[0] or "") and str(r[1] or "").strip() == "Data Item"
        ),
        None,
    )
    if head_at is None:
        yield Rejected(path.name, "no 'NAICS Code' / 'Data Item' header row")
        return
    years = _years(rows[head_at])
    if not years:
        yield Rejected(path.name, "no year columns in the header row")
        return
    found: set[str] = set()
    for i, r in enumerate(rows[head_at + 1 :], start=head_at + 2):
        if not r or len(r) < 4 or str(r[1] or "").strip() != AWTS_ITEM:
            continue
        code = str(r[0]).strip() if r[0] is not None else ""
        if code not in AWTS:
            continue
        measure, want = AWTS[code]
        label = _label(r[2])
        if label != want or str(r[3] or "").strip() != AWTS_OPERATION or code in found:
            yield Rejected(
                f"{path.name} row {i}",
                f"NAICS {code}: {label!r} / {str(r[3] or '').strip()!r}, expected {want!r} / {AWTS_OPERATION!r}, or a second row",
                measure_id=measure,
            )
            continue
        found.add(code)
        yield from _cells(asset, path, ws.title, i, r, years, measure, code, label, asset.source_series)
    for code in AWTS.keys() - found:
        yield Rejected(path.name, f"NAICS {code} ({AWTS[code][1]}) not found", measure_id=AWTS[code][0])
